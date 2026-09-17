import argparse
import hashlib
import json
import resource
import sqlite3
import sys
import tempfile
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime, timezone
from itertools import combinations
from pathlib import Path

import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook

from h1b_job_insights import schema
from h1b_job_insights.sources import Source, discover_sources

REQUIRED = {
    "main": {
        "CASE_NUMBER",
        "CASE_STATUS",
        "DECISION_DATE",
        "RECEIVED_DATE",
        "ORIGINAL_CERT_DATE",
        "VISA_CLASS",
        "EMPLOYER_NAME",
        "SOC_CODE",
        "WORKSITE_CITY",
        "WORKSITE_STATE",
        "WORKSITE_POSTAL_CODE",
    },
    "worksites": {
        "CASE_NUMBER",
        "WORKSITE_CITY",
        "WORKSITE_STATE",
        "WORKSITE_POSTAL_CODE",
        "WORKSITE_WORKERS",
    },
}
PROVENANCE = ("SOURCE_RELEASE", "SOURCE_FILE", "SOURCE_ROW_NUMBER")
START_DATE = date(2022, 1, 1)
BATCH_SIZE = 10_000
STATE_CODES = {
    "ALABAMA": "AL",
    "ALASKA": "AK",
    "ARIZONA": "AZ",
    "ARKANSAS": "AR",
    "CALIFORNIA": "CA",
    "COLORADO": "CO",
    "CONNECTICUT": "CT",
    "DELAWARE": "DE",
    "DISTRICT OF COLUMBIA": "DC",
    "FLORIDA": "FL",
    "GEORGIA": "GA",
    "GUAM": "GU",
    "HAWAII": "HI",
    "IDAHO": "ID",
    "ILLINOIS": "IL",
    "INDIANA": "IN",
    "IOWA": "IA",
    "KANSAS": "KS",
    "KENTUCKY": "KY",
    "LOUISIANA": "LA",
    "MAINE": "ME",
    "MARYLAND": "MD",
    "MASSACHUSETTS": "MA",
    "MICHIGAN": "MI",
    "MINNESOTA": "MN",
    "MISSISSIPPI": "MS",
    "MISSOURI": "MO",
    "MONTANA": "MT",
    "NEBRASKA": "NE",
    "NEVADA": "NV",
    "NEW HAMPSHIRE": "NH",
    "NEW JERSEY": "NJ",
    "NEW MEXICO": "NM",
    "NEW YORK": "NY",
    "NORTH CAROLINA": "NC",
    "NORTH DAKOTA": "ND",
    "NORTHERN MARIANA ISLANDS": "MP",
    "OHIO": "OH",
    "OKLAHOMA": "OK",
    "OREGON": "OR",
    "PENNSYLVANIA": "PA",
    "PUERTO RICO": "PR",
    "RHODE ISLAND": "RI",
    "SOUTH CAROLINA": "SC",
    "SOUTH DAKOTA": "SD",
    "TENNESSEE": "TN",
    "TEXAS": "TX",
    "UTAH": "UT",
    "VERMONT": "VT",
    "VIRGIN ISLANDS": "VI",
    "VIRGINIA": "VA",
    "WASHINGTON": "WA",
    "WEST VIRGINIA": "WV",
    "WISCONSIN": "WI",
    "WYOMING": "WY",
}


@dataclass(frozen=True)
class WorkbookInfo:
    source: Source
    path: Path
    sheet: str
    headers: tuple[str, ...]
    mapped_headers: tuple[str, ...]
    sha256: str
    size_bytes: int
    local_modified_at: str


def cell_text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        if value.time().isoformat() == "00:00:00":
            return value.date().isoformat()
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def decision_date(value: str | None) -> date | None:
    if not value:
        return None
    try:
        return date.fromisoformat(value[:10])
    except ValueError:
        return None


def location(record: dict[str, str | None]) -> tuple[str, str, str]:
    city = (record.get("WORKSITE_CITY") or "").strip().upper()
    state = (record.get("WORKSITE_STATE") or "").strip().upper()
    postal = (record.get("WORKSITE_POSTAL_CODE") or "").strip().upper()
    if postal.isdigit() and len(postal) < 5:
        postal = postal.zfill(5)
    return city, STATE_CODES.get(state, state), postal


def inspect(source: Source, path: Path) -> WorkbookInfo:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(8 * 1024 * 1024):
            digest.update(chunk)

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.sheetnames) != 1:
            raise ValueError(f"Expected one sheet in {path}, found {workbook.sheetnames}")
        sheet = workbook.active
        headers = tuple(sheet.cell(1, col).value for col in range(1, sheet.max_column + 1))
        if any(not isinstance(header, str) or not header for header in headers):
            raise ValueError(f"Blank or non-text header in {path}")
        if len(headers) != len(set(headers)):
            raise ValueError(f"Duplicate headers in {path}")
        mapped_headers = schema.map_headers(source.kind, headers, path)
        missing = REQUIRED[source.kind] - set(mapped_headers)
        if missing:
            raise ValueError(f"Missing required headers in {path}: {', '.join(sorted(missing))}")
        if set(headers) & set(PROVENANCE):
            raise ValueError(f"Source header conflicts with provenance columns in {path}")
        return WorkbookInfo(
            source=source,
            path=path,
            sheet=sheet.title,
            headers=headers,
            mapped_headers=mapped_headers,
            sha256=digest.hexdigest(),
            size_bytes=path.stat().st_size,
            local_modified_at=datetime.fromtimestamp(
                path.stat().st_mtime, timezone.utc
            ).isoformat(),
        )
    finally:
        workbook.close()


def convert(
    info: WorkbookInfo,
    columns: tuple[str, ...],
    output_dir: Path,
    on_row,
) -> int:
    target = output_dir / info.source.kind / f"{info.source.release.lower()}.parquet"
    target.parent.mkdir(parents=True, exist_ok=True)
    partial = target.with_suffix(".parquet.part")
    schema = pa.schema(
        [(column, pa.string()) for column in columns]
        + [
            ("SOURCE_RELEASE", pa.string()),
            ("SOURCE_FILE", pa.string()),
            ("SOURCE_ROW_NUMBER", pa.int64()),
        ]
    )
    mapped_headers = info.mapped_headers
    absent = tuple(column for column in columns if column not in mapped_headers)
    buffer = {field.name: [] for field in schema}
    workbook = load_workbook(info.path, read_only=True, data_only=True)
    rows = 0
    try:
        sheet = workbook.active
        actual_headers = tuple(next(sheet.iter_rows(values_only=True)))
        if actual_headers != info.headers:
            raise ValueError(f"Headers changed while reading {info.path}")
        with pq.ParquetWriter(partial, schema, compression="zstd") as writer:
            for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                if not any(value is not None for value in row):
                    continue
                if any(value is not None for value in row[len(info.headers) :]):
                    raise ValueError(
                        f"Unexpected data beyond headers in {info.path}, row {row_number}"
                    )
                values = [cell_text(value) for value in row[: len(info.headers)]]
                values.extend([None] * (len(info.headers) - len(values)))
                record = dict(zip(mapped_headers, values))
                on_row(record)
                for header, value in record.items():
                    buffer[header].append(value)
                for header in absent:
                    buffer[header].append(None)
                buffer["SOURCE_RELEASE"].append(info.source.release)
                buffer["SOURCE_FILE"].append(info.source.filename)
                buffer["SOURCE_ROW_NUMBER"].append(row_number)
                rows += 1
                if rows % BATCH_SIZE == 0:
                    writer.write_table(pa.Table.from_pydict(buffer, schema=schema))
                    buffer = {field.name: [] for field in schema}
            if buffer["SOURCE_ROW_NUMBER"]:
                writer.write_table(pa.Table.from_pydict(buffer, schema=schema))
        partial.replace(target)
        return rows
    except Exception:
        partial.unlink(missing_ok=True)
        raise
    finally:
        workbook.close()


def add_main(
    info: WorkbookInfo, columns: tuple[str, ...], output_dir: Path, db: sqlite3.Connection
) -> dict:
    stats = {
        "release": info.source.release,
        "visa_class_counts": Counter(),
        "case_status_counts": Counter(),
        "h1b_case_status_counts": Counter(),
        "h1b_rows": 0,
        "h1b_rows_since_2022_01_01": 0,
        "missing_decision_date": 0,
        "invalid_decision_date": 0,
        "h1b_missing_decision_date": 0,
        "h1b_missing_employer_name": 0,
        "h1b_missing_soc_code": 0,
        "missing_case_number": 0,
        "duplicate_case_number_rows": 0,
    }
    pending = []
    earliest = latest = None

    def flush_cases() -> None:
        if not pending:
            return
        before = db.total_changes
        db.executemany("INSERT OR IGNORE INTO cases VALUES (?, ?, ?, ?, ?, ?)", pending)
        stats["duplicate_case_number_rows"] += len(pending) - (db.total_changes - before)
        pending.clear()

    def on_row(record: dict[str, str | None]) -> None:
        nonlocal earliest, latest
        visa = record["VISA_CLASS"] or "<missing>"
        status = record["CASE_STATUS"] or "<missing>"
        stats["visa_class_counts"][visa] += 1
        stats["case_status_counts"][status] += 1
        h1b = visa.strip().upper() == "H-1B"
        if h1b:
            stats["h1b_rows"] += 1
            stats["h1b_case_status_counts"][status] += 1
            if not record["EMPLOYER_NAME"]:
                stats["h1b_missing_employer_name"] += 1
            if not record["SOC_CODE"]:
                stats["h1b_missing_soc_code"] += 1
        raw_date = record["DECISION_DATE"]
        parsed = decision_date(raw_date)
        if not raw_date:
            stats["missing_decision_date"] += 1
            if h1b:
                stats["h1b_missing_decision_date"] += 1
        elif parsed is None:
            stats["invalid_decision_date"] += 1
        else:
            earliest = min(earliest, parsed) if earliest else parsed
            latest = max(latest, parsed) if latest else parsed
            if h1b and parsed >= START_DATE:
                stats["h1b_rows_since_2022_01_01"] += 1
        case_number = record["CASE_NUMBER"]
        if not case_number:
            stats["missing_case_number"] += 1
            return
        pending.append(
            (
                info.source.release,
                info.source.release.split("_")[0],
                case_number,
                *location(record),
            )
        )
        if len(pending) == BATCH_SIZE:
            flush_cases()

    stats["source_rows"] = convert(info, columns, output_dir, on_row)
    flush_cases()
    stats["unique_case_numbers"] = (
        stats["source_rows"] - stats["missing_case_number"] - stats["duplicate_case_number_rows"]
    )
    stats["decision_date_min"] = earliest.isoformat() if earliest else None
    stats["decision_date_max"] = latest.isoformat() if latest else None
    return stats


def add_worksites(
    info: WorkbookInfo, columns: tuple[str, ...], output_dir: Path, db: sqlite3.Connection,
    *, reuse: bool = False,
) -> dict:
    year = info.source.release.split("_")[0]
    pending = []
    seen = set()
    multiple = set()
    matched = set()
    comparable = set()
    first_matches = set()
    matched_rows = 0
    missing_case = 0

    def flush_lookups() -> None:
        nonlocal matched_rows
        if not pending:
            return
        case_numbers = list({case for case, _ in pending})
        main_locations: dict[str, list[tuple[str, str, str]]] = {}
        for start in range(0, len(case_numbers), 500):
            part = case_numbers[start : start + 500]
            placeholders = ",".join("?" for _ in part)
            query = (
                "SELECT case_number, city, state, postal FROM cases "
                f"WHERE fiscal_year = ? AND case_number IN ({placeholders})"
            )
            for case, city, state, postal in db.execute(query, [year, *part]):
                main_locations.setdefault(case, []).append((city, state, postal))
        for case, worksite_location in pending:
            locations = main_locations.get(case)
            if not locations:
                continue
            matched_rows += 1
            matched.add(case)
            if any(all(value for value in loc) for loc in locations):
                comparable.add(case)
            if all(worksite_location) and worksite_location in locations:
                first_matches.add(case)
        pending.clear()

    def on_row(record: dict[str, str | None]) -> None:
        nonlocal missing_case
        case = record["CASE_NUMBER"]
        if not case:
            missing_case += 1
            return
        if case in seen:
            multiple.add(case)
        else:
            seen.add(case)
        pending.append((case, location(record)))
        if len(pending) == BATCH_SIZE:
            flush_lookups()

    if reuse:
        path = output_dir / "worksites" / f"{info.source.release.lower()}.parquet"
        source_rows = pq.read_metadata(path).num_rows
        fields = ["CASE_NUMBER", "WORKSITE_CITY", "WORKSITE_STATE", "WORKSITE_POSTAL_CODE"]
        for batch in pq.ParquetFile(path).iter_batches(columns=fields, batch_size=BATCH_SIZE):
            for record in batch.to_pylist():
                on_row(record)
    else:
        source_rows = convert(info, columns, output_dir, on_row)
    flush_lookups()
    unique_main = db.execute(
        "SELECT COUNT(DISTINCT case_number) FROM cases WHERE fiscal_year = ?", (year,)
    ).fetchone()[0]
    return {
        "release": info.source.release,
        "source_rows": source_rows,
        "missing_case_number": missing_case,
        "unique_worksite_case_numbers": len(seen),
        "cases_with_multiple_worksite_rows": len(multiple),
        "worksite_rows_matching_main": matched_rows,
        "worksite_rows_without_main": source_rows - missing_case - matched_rows,
        "unique_main_case_numbers": unique_main,
        "main_cases_with_worksite_row": len(matched),
        "main_cases_without_worksite_row": unique_main - len(matched),
        "main_to_worksite_match_pct": round(100 * len(matched) / unique_main, 2)
        if unique_main
        else None,
        "worksite_to_main_match_pct": round(100 * matched_rows / (source_rows - missing_case), 2)
        if source_rows > missing_case
        else None,
        "main_cases_with_comparable_first_worksite": len(comparable),
        "main_cases_with_first_worksite_location_match": len(first_matches),
        "first_worksite_location_match_pct": round(100 * len(first_matches) / len(comparable), 2)
        if comparable
        else None,
    }


def write_json(path: Path, value: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    partial = path.with_suffix(path.suffix + ".part")
    partial.write_text(json.dumps(value, indent=2, ensure_ascii=False) + "\n")
    partial.replace(path)


def audit_case_links(report: dict, output_dir: Path) -> None:
    main_by_release = {}
    for item in report["main"]:
        release = item["release"]
        path = output_dir / "main" / f"{release.lower()}.parquet"
        main_by_release[release] = {
            case
            for case in pq.read_table(path, columns=["CASE_NUMBER"])["CASE_NUMBER"].to_pylist()
            if case is not None
        }

    overlaps = []
    for (left, left_cases), (right, right_cases) in combinations(main_by_release.items(), 2):
        count = len(left_cases & right_cases)
        if count:
            overlaps.append({"left": left, "right": right, "shared_case_numbers": count})
    report["overlapping_releases"] = overlaps
    all_main = set().union(*main_by_release.values())
    main_by_year = {}
    for release, cases in main_by_release.items():
        main_by_year.setdefault(release.split("_")[0], set()).update(cases)
    del main_by_release

    all_worksites = set()
    for item in report["worksites"]:
        release = item["release"]
        year = release.split("_")[0]
        path = output_dir / "worksites" / f"{release.lower()}.parquet"
        case_numbers = pq.read_table(path, columns=["CASE_NUMBER"])["CASE_NUMBER"].to_pylist()
        worksite_cases = {case for case in case_numbers if case is not None}
        all_worksites.update(worksite_cases)
        renamed = {
            "worksite_rows_without_main": "worksite_rows_without_same_year_main",
            "main_cases_with_worksite_row": "main_cases_with_same_year_worksite_row",
            "main_cases_without_worksite_row": "main_cases_without_same_year_worksite_row",
            "main_to_worksite_match_pct": "same_year_main_to_worksite_match_pct",
            "worksite_to_main_match_pct": "same_year_worksite_to_main_match_pct",
        }
        for old, new in renamed.items():
            if old in item:
                item[new] = item.pop(old)
        item["worksite_cases_with_main_in_other_year"] = len(
            (worksite_cases - main_by_year.get(year, set())) & all_main
        )
        item["worksite_cases_without_any_selected_main"] = len(worksite_cases - all_main)
        item["worksite_rows_without_any_selected_main"] = sum(
            case not in all_main for case in case_numbers if case is not None
        )

    matched = len(all_main & all_worksites)
    report["global_case_links_by_fiscal_year"] = [
        {
            "fiscal_year": year,
            "unique_main_case_numbers": len(cases),
            "main_cases_with_any_worksite_row": len(cases & all_worksites),
            "main_cases_without_any_worksite_row": len(cases - all_worksites),
            "main_to_worksite_match_pct": round(100 * len(cases & all_worksites) / len(cases), 2)
            if cases
            else None,
        }
        for year, cases in sorted(main_by_year.items())
    ]
    report["global_case_links"] = {
        "unique_main_case_numbers": len(all_main),
        "unique_worksite_case_numbers": len(all_worksites),
        "main_cases_with_any_worksite_row": matched,
        "main_cases_without_any_worksite_row": len(all_main) - matched,
        "main_to_worksite_match_pct": round(100 * matched / len(all_main), 2) if all_main else None,
        "worksite_cases_without_any_selected_main": len(all_worksites - all_main),
    }


def processor_sha256() -> str:
    digest = hashlib.sha256()
    for path in (Path(__file__), Path(schema.__file__)):
        digest.update(path.read_bytes())
    return digest.hexdigest()


def manifest_for(
    infos: list[WorkbookInfo], row_counts: dict[str, int],
    main_columns: tuple[str, ...], worksite_columns: tuple[str, ...],
) -> dict:
    return {
        "processor_sha256": processor_sha256(),
        "files": [
            {
                "release": info.source.release,
                "kind": info.source.kind,
                "url": info.source.url,
                "filename": info.source.filename,
                "local_modified_at_utc": info.local_modified_at,
                "sha256": info.sha256,
                "size_bytes": info.size_bytes,
                "sheets": [
                    {
                        "name": info.sheet,
                        "data_rows": row_counts[info.source.filename],
                        "headers": list(info.headers),
                        "mapped_headers": list(info.mapped_headers),
                    }
                ],
            }
            for info in infos
        ],
        "processed_columns": {
            "main": list(main_columns) + list(PROVENANCE),
            "worksites": list(worksite_columns) + list(PROVENANCE),
        },
    }


def cached_report(
    raw_dir: Path, output_dir: Path, sources: tuple[Source, ...]
) -> dict | None:
    manifest_path = output_dir / "source_manifest.json"
    quality_path = output_dir / "quality.json"
    if not manifest_path.is_file() or not quality_path.is_file():
        return None
    try:
        manifest = json.loads(manifest_path.read_text())
        previous = manifest["files"]
        if len(previous) != len(sources):
            return None
        if manifest.get("processor_sha256") != processor_sha256():
            return None
        for source, saved in zip(sources, previous):
            file_stat = (raw_dir / source.filename).stat()
            if any(
                saved.get(key) != value
                for key, value in {
                    "release": source.release,
                    "kind": source.kind,
                    "url": source.url,
                    "filename": source.filename,
                    "size_bytes": file_stat.st_size,
                    "local_modified_at_utc": datetime.fromtimestamp(
                        file_stat.st_mtime, timezone.utc
                    ).isoformat(),
                }.items()
            ):
                return None
            sheet = saved["sheets"][0]
            path = output_dir / source.kind / f"{source.release.lower()}.parquet"
            if not path.is_file() or pq.read_metadata(path).num_rows != sheet["data_rows"]:
                return None
        return json.loads(quality_path.read_text())
    except (KeyError, IndexError, OSError, ValueError, TypeError):
        return None


def previous_results(output_dir: Path) -> tuple[dict, dict]:
    try:
        manifest = json.loads((output_dir / "source_manifest.json").read_text())
        if manifest.get("processor_sha256") != processor_sha256():
            return {}, {}
        report = json.loads((output_dir / "quality.json").read_text())
        return manifest, report
    except (OSError, ValueError):
        return {}, {}


def reusable_info(
    source: Source, raw_dir: Path, output_dir: Path, saved: dict | None
) -> WorkbookInfo | None:
    if not saved:
        return None
    path = raw_dir / source.filename
    file_stat = path.stat()
    expected = {
        "release": source.release,
        "kind": source.kind,
        "url": source.url,
        "filename": source.filename,
        "size_bytes": file_stat.st_size,
        "local_modified_at_utc": datetime.fromtimestamp(
            file_stat.st_mtime, timezone.utc
        ).isoformat(),
    }
    if any(saved.get(key) != value for key, value in expected.items()):
        return None
    try:
        sheet = saved["sheets"][0]
        headers = tuple(sheet["headers"])
        mapped = schema.map_headers(source.kind, headers, path)
        output = output_dir / source.kind / f"{source.release.lower()}.parquet"
        if pq.read_metadata(output).num_rows != sheet["data_rows"]:
            return None
        if pq.read_schema(output).names != list(schema.COLUMNS[source.kind]) + list(PROVENANCE):
            return None
        return WorkbookInfo(
            source=source,
            path=path,
            sheet=sheet["name"],
            headers=headers,
            mapped_headers=mapped,
            sha256=saved["sha256"],
            size_bytes=file_stat.st_size,
            local_modified_at=expected["local_modified_at_utc"],
        )
    except (KeyError, IndexError, OSError, ValueError, TypeError):
        return None


def index_main(path: Path, release: str, db: sqlite3.Connection) -> None:
    fields = ["CASE_NUMBER", "WORKSITE_CITY", "WORKSITE_STATE", "WORKSITE_POSTAL_CODE"]
    year = release.split("_")[0]
    for batch in pq.ParquetFile(path).iter_batches(columns=fields, batch_size=BATCH_SIZE):
        rows = []
        for record in batch.to_pylist():
            if case := record["CASE_NUMBER"]:
                rows.append((release, year, case, *location(record)))
        db.executemany("INSERT OR IGNORE INTO cases VALUES (?, ?, ?, ?, ?, ?)", rows)


def run(
    raw_dir: Path,
    output_dir: Path,
    sources: tuple[Source, ...] | None = None,
    *,
    force: bool = False,
) -> dict:
    if sources is None:
        sources = discover_sources(raw_dir)
    missing = [
        raw_dir / source.filename
        for source in sources
        if not (raw_dir / source.filename).is_file()
    ]
    if missing:
        raise FileNotFoundError(
            "Missing source workbooks. Save these files before running the pipeline:\n"
            + "\n".join(str(path) for path in missing)
        )
    if not force and (report := cached_report(raw_dir, output_dir, sources)) is not None:
        print("Source files unchanged; using existing processed data", flush=True)
        return report
    old_manifest, old_report = previous_results(output_dir) if not force else ({}, {})
    old_files = {item["filename"]: item for item in old_manifest.get("files", [])}
    old_main = {item["release"]: item for item in old_report.get("main", [])}
    infos = []
    reused = set()
    for source in sources:
        info = reusable_info(source, raw_dir, output_dir, old_files.get(source.filename))
        if info is not None and (source.kind != "main" or source.release in old_main):
            reused.add(source.filename)
        else:
            print(f"Inspecting {source.release} {source.kind}", flush=True)
            info = inspect(source, raw_dir / source.filename)
        infos.append(info)
    main_columns = schema.MAIN_COLUMNS
    worksite_columns = schema.WORKSITE_COLUMNS
    output_dir.mkdir(parents=True, exist_ok=True)
    report = {"main": [], "worksites": []}
    row_counts = {}
    with tempfile.TemporaryDirectory(prefix="validation-", dir=output_dir) as temp:
        db = sqlite3.connect(Path(temp) / "cases.sqlite")
        db.execute("PRAGMA journal_mode=OFF")
        db.execute("PRAGMA synchronous=OFF")
        db.execute("PRAGMA cache_size=-64000")
        db.execute(
            "CREATE TABLE cases (release TEXT, fiscal_year TEXT, case_number TEXT, "
            "city TEXT, state TEXT, postal TEXT, PRIMARY KEY (release, case_number)) WITHOUT ROWID"
        )
        for info in infos:
            if info.source.kind != "main":
                continue
            if info.source.filename in reused:
                print(f"Reusing {info.source.release} main", flush=True)
                path = output_dir / "main" / f"{info.source.release.lower()}.parquet"
                index_main(path, info.source.release, db)
                stats = old_main[info.source.release]
            else:
                print(f"Reading {info.source.release} main", flush=True)
                stats = add_main(info, main_columns, output_dir, db)
            report["main"].append(stats)
            row_counts[info.source.filename] = stats["source_rows"]
            print(f"  {stats['source_rows']:,} rows", flush=True)
        db.execute("CREATE INDEX cases_by_year ON cases (fiscal_year, case_number)")
        duplicated = db.execute(
            "SELECT COUNT(*), COALESCE(SUM(n - 1), 0) FROM "
            "(SELECT COUNT(*) AS n FROM cases GROUP BY case_number HAVING n > 1)"
        ).fetchone()
        report["case_numbers_in_multiple_releases"] = duplicated[0]
        report["extra_case_occurrences_across_releases"] = duplicated[1]
        for info in infos:
            if info.source.kind != "worksites":
                continue
            reuse = info.source.filename in reused
            action = "Reusing" if reuse else "Reading"
            print(f"{action} {info.source.release} worksites", flush=True)
            stats = add_worksites(info, worksite_columns, output_dir, db, reuse=reuse)
            report["worksites"].append(stats)
            row_counts[info.source.filename] = stats["source_rows"]
            print(f"  {stats['source_rows']:,} rows", flush=True)
        db.close()

    audit_case_links(report, output_dir)
    manifest = manifest_for(infos, row_counts, main_columns, worksite_columns)
    peak = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    report["peak_rss_mib"] = round(peak / (1024 * 1024 if sys.platform == "darwin" else 1024), 1)
    write_json(output_dir / "source_manifest.json", manifest)
    write_json(output_dir / "quality.json", report)
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description="Process locally saved DOL LCA workbooks")
    parser.add_argument("--raw-dir", type=Path, default=Path("data/raw"))
    parser.add_argument("--output-dir", type=Path, default=Path("data/processed"))
    parser.add_argument("--force", action="store_true", help="Rebuild even if inputs are unchanged")
    args = parser.parse_args()
    run(args.raw_dir, args.output_dir, force=args.force)
    print(f"Results: {args.output_dir / 'source_manifest.json'} and quality.json")


if __name__ == "__main__":
    main()
